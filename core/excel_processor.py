import pandas as pd
import openpyxl
from config.settings import EXCEL

class ExcelProcessor:
    @staticmethod
    def write_results_to_xlsm(excel_path, results: list[dict], path_col_name: str):
        """
        Menulis hasil judgement langsung ke file .xlsm asli menggunakan openpyxl,
        sehingga makro VBA tidak rusak.

        Args:
        excel_path  : Path ke file .xlsm asli.
        results     : List of dict, setiap dict berisi:
                        { 'depot_path', 'necessity', 'impl_or_not', 'reason' }
        path_col_name: Nama kolom File Path di Excel.
    """
    wb = openpyxl.load_workbook(excel_path, keep_vba=True)

    if EXCEL.DEFAULT_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet '{EXCEL.DEFAULT_SHEET_NAME}' tidak ditemukan!")

    ws = wb[EXCEL.DEFAULT_SHEET_NAME]

    # --- Temukan baris header dan index kolom yang dibutuhkan ---
    header_row_idx  = None
    col_path        = None
    col_necessity   = None
    col_impl        = None
    col_reason      = None

    for row in ws.iter_rows():
        for cell in row:
            if str(cell.value).strip() == path_col_name:
                header_row_idx = cell.row
                break
        if header_row_idx:
            break

    if not header_row_idx:
        raise ValueError(f"Kolom '{path_col_name}' tidak ditemukan di sheet!")

    # Baca nama kolom dari baris header
    for cell in ws[header_row_idx]:
        val = str(cell.value).strip() if cell.value else ""
        if val == path_col_name:
            col_path = cell.column
        elif val == "Necessity of Evaluation":
            col_necessity = cell.column
        elif val == "Implementation or Not":
            col_impl = cell.column
        elif val == "Reason":
            col_reason = cell.column

    # Buat lookup dict: depot_path → result
    result_map = {r['depot_path']: r for r in results}

    # --- Tulis hasil ke setiap baris data ---
    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        path_cell = ws.cell(row=row_idx, column=col_path)
        if path_cell.value is None:
            continue

        depot_path = str(path_cell.value).strip()
        if depot_path not in result_map:
            continue

        res = result_map[depot_path]

        if col_necessity:
            ws.cell(row=row_idx, column=col_necessity).value = res['necessity']
        if col_impl:
            ws.cell(row=row_idx, column=col_impl).value = res['impl_or_not']
        if col_reason:
            ws.cell(row=row_idx, column=col_reason).value = res['reason']

    # Simpan kembali ke file xlsm ASLI
    wb.save(excel_path)

    @staticmethod
    def load_and_prepare(excel_path, path_col_name):
        """Membaca file Excel, mencari baris header yang tepat, dan menyiapkan kolom output."""
        try:
            xl_file = pd.ExcelFile(excel_path)
            if EXCEL.DEFAULT_SHEET_NAME not in xl_file.sheet_names:
                return None, f"Sheet '{EXCEL.DEFAULT_SHEET_NAME}' tidak ditemukan!"

            df_temp = pd.read_excel(excel_path, sheet_name=EXCEL.DEFAULT_SHEET_NAME, header=None, nrows=EXCEL.MAX_HEADER_ROWS_TO_CHECK)
            header_row_index = -1
            
            for idx, row in df_temp.iterrows():
                row_values = [str(x).strip() for x in row.values]
                if path_col_name in row_values:
                    header_row_index = idx
                    break
            
            if header_row_index == -1:
                return None, f"Kolom '{path_col_name}' tidak ditemukan di {EXCEL.MAX_HEADER_ROWS_TO_CHECK} baris pertama!"
            
            df = pd.read_excel(excel_path, sheet_name=EXCEL.DEFAULT_SHEET_NAME, header=header_row_index)
            df.columns = [str(c).strip() for c in df.columns]
            
            if path_col_name not in df.columns:
                 return None, f"Kolom '{path_col_name}' gagal dibaca."
            
            if 'Target CL' not in df.columns:
                return None, "Kolom 'Target CL' tidak ditemukan di Excel."

            if 'Necessity of Evaluation' not in df.columns:
                df['Necessity of Evaluation'] = ""
            if 'Reason' not in df.columns:
                df['Reason'] = ""

            return df, "OK"
        except Exception as e:
            return None, str(e)

    @staticmethod
    def save(df, output_path):
        """Menyimpan DataFrame kembali ke Excel."""
        df.to_excel(output_path, index=False)